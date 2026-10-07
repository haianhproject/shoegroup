from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor

out = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.1
doc.styles['Title'].font.size = Pt(18)
doc.styles['Heading 1'].font.size = Pt(12)
doc.add_paragraph('Phân công công việc ShoeGroup', 'Title')
doc.add_paragraph('Ngày 7/10/26 | Hải Anh nhóm trưởng | Kế hoạch tối nay')
tasks = [
('Hải Anh', 'fix menu thuộc tính sản phẩm, kiểm tra màu và size', '3h', '19:00–22:00', 'Products, ProductVariants, ProductImages, Categories, Brands, Materials, Colors, Sizes.'),
('Hưng', 'kiểm tra tạo đơn, thanh toán và cập nhật trạng thái đơn', '3h', '19:00–22:00', 'Orders, OrderDetails, OrderStatusHistory, PaymentTransactions, ShippingMethods, CheckoutRequests.'),
('Hiếu', 'kiểm tra đăng nhập, phân quyền, địa chỉ và thông báo', '2h', '19:00–21:00', 'Users, Roles, UserAddresses, Notifications.'),
('Việt', 'kiểm tra giỏ online/POS, tồn kho, mã giảm giá và ưu đãi', '3h', '19:00–22:00', 'Carts, CartItems, PosCarts, PosCartItems, Coupons, CouponRedemptions, VariantDiscounts.'),
]
for name, task, duration, hours, tables in tasks:
    p = doc.add_paragraph()
    p.add_run(f'Ngày 7/10/26: {name} ').bold = True
    p.add_run(f'{task} (dự kiến {duration} xong, {hours}).')
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph('Bảng phụ trách: ' + tables)
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(10)
doc.add_paragraph('Báo cáo hôm qua và giao việc buổi sáng', 'Heading 1')
doc.add_paragraph('Ngày 6/10/26: Hải Anh, Hưng, Hiếu, Việt chưa có báo cáo.')
doc.add_paragraph('Mỗi sáng 09:00: từng người báo cáo hôm qua làm được gì, hôm nay sẽ làm gì; Hải Anh tổng hợp và giao task. Tối 22:30 chốt kết quả thực tế.')
doc.add_paragraph('Mẫu: Ngày …: [Tên] — Hôm qua: …; Hôm nay: … (dự kiến … giờ xong, lúc …); Vướng mắc: …')
doc.add_paragraph('Bảng cần rà soát', 'Heading 1')
doc.add_paragraph('Việt rà Carts/CartItems còn API cũ; Hiếu rà Roles chưa có truy vấn danh mục trực tiếp nhưng còn liên kết Users.RoleID. Chưa đủ căn cứ kết luận không dùng hoặc xóa.')
doc.core_properties.author = ''
doc.core_properties.title = 'Phân công công việc ShoeGroup ngày 7 tháng 10 năm 2026'
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter()):
        if border.tag.endswith('}pBdr'):
            border.getparent().remove(border)
path = out / 'Phan-cong-ShoeGroup-07-10-2026.docx'
doc.save(path)
print(path)
